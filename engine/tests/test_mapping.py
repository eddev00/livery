"""Import mapping: settings and rules (docs/MAPPING.md)."""
from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from docx.enum.text import WD_BREAK

from livery import pipeline
from livery.convert import DEFAULT_THEME, convert, load_theme
from livery.mapping import MappingError, apply_to_theme, fingerprint, parse


def _mapping_doc(path: Path) -> Path:
    """Cover page, a leftover page, then the body: bold numbered titles typed in Normal
    paragraphs, a custom title style, and a table used as a container for another table."""
    doc = Document()
    doc.add_paragraph("Annual Plan")
    doc.add_paragraph("Prepared for the board").runs[0].add_break(WD_BREAK.PAGE)
    doc.add_paragraph("Internal distribution only")
    doc.add_paragraph("Do not circulate").runs[0].add_break(WD_BREAK.PAGE)
    doc.add_paragraph().add_run("1. Introduction").bold = True
    doc.add_paragraph("The plan in short.")
    doc.add_paragraph().add_run("1.1 Scope").bold = True
    doc.add_paragraph("What the plan covers.")
    doc.styles.add_style("Body Title", 1)
    doc.add_paragraph("Budget overview", style="Body Title")
    outer = doc.add_table(rows=1, cols=1)
    outer.style = "Table Grid"
    cell = outer.rows[0].cells[0]
    cell.paragraphs[0].text = "Figures"
    inner = cell.add_table(rows=2, cols=2)
    for r in range(2):
        for c in range(2):
            inner.rows[r].cells[c].text = f"Value {r}{c}"
    doc.add_paragraph("Closing words.")
    doc.save(path)
    return path


RULES = """
[settings]
headings = { "Body Title" = 2 }

[[rule]]
name = "Leftover page"
match = { page = 2 }
action = "drop"
expect = "1+"

[[rule]]
name = "Bold numbered titles"
match = { kind = "paragraph", bold = true, numbered = true }
action = { heading = "from_number" }

[[rule]]
name = "Container tables"
match = { kind = "table", contains_table = true }
action = { nested_table = "unwrap_outer" }
"""


@pytest.fixture(scope="module")
def converted(tmp_path_factory):
    d = tmp_path_factory.mktemp("mapping")
    src = _mapping_doc(d / "plan.docx")
    theme = apply_to_theme(load_theme(DEFAULT_THEME), parse(RULES))
    return theme, convert(src, d / "out", theme)


def _headings(conv) -> list[tuple[int, str]]:
    return [(b["c"][0], pipeline.past.stringify_inlines(b["c"][2]))
            for b in conv.source_ast["blocks"] if b["t"] == "Header"]


def test_rules_restructure_the_document(converted):
    _, conv = converted
    assert conv.rule_counts == {"Leftover page": 2, "Bold numbered titles": 2, "Container tables": 1}
    # the typed numbers became the levels; the template regenerates them (heading_numbers = "parse")
    assert _headings(conv) == [(1, "Introduction"), (2, "Scope"), (2, "Budget overview")]
    # the outer table is gone, the inner one is now at top level
    assert conv.structure.tables == 1 and "Value 11" in conv.source_text
    assert "Internal distribution" not in conv.source_text and "Annual Plan" in conv.source_text


def test_removals_are_declared_and_every_word_accounted_for(converted):
    theme, conv = converted
    assert conv.dropped_paragraphs == ["[rule:Leftover page] Internal distribution only",
                                       "[rule:Leftover page] Do not circulate"]
    rules = pipeline._rules_check(theme["rule"], conv.rule_counts, conv.dropped_paragraphs, conv.raw_text)
    assert rules.status == "PASS" and any("Internal distribution only" in d for d in rules.details)
    text = pipeline._mapping_text_check(conv.raw_text, conv.source_text, conv.dropped_paragraphs, conv.headings)
    assert text.status == "PASS", text.details


def test_text_lost_without_a_record_fails(converted):
    _, conv = converted
    lost = conv.source_text.replace("What the plan covers.", "")
    check = pipeline._mapping_text_check(conv.raw_text, lost, conv.dropped_paragraphs, conv.headings)
    assert check.status == "FAIL" and "covers" in check.details[0]


def test_runaway_and_unexpected_rules_are_flagged(converted):
    theme, conv = converted
    rules = [dict(r) for r in theme["rule"]]
    rules[0]["expect"] = "0"
    assert pipeline._rules_check(rules, conv.rule_counts, conv.dropped_paragraphs, conv.raw_text).status == "WARN"
    huge = conv.dropped_paragraphs + [f"[rule:Leftover page] {conv.source_text}"]
    assert pipeline._rules_check(rules, conv.rule_counts, huge, conv.raw_text).status == "FAIL"


def test_keep_overrides_a_built_in_drop(tmp_path):
    doc = Document()
    doc.add_heading("Contents", level=1)
    doc.add_paragraph("Body.")
    src = tmp_path / "c.docx"
    doc.save(src)
    base = load_theme(DEFAULT_THEME)
    plain = convert(src.resolve(), tmp_path / "a", base)
    assert "Contents" not in plain.source_text
    kept = convert(src.resolve(), tmp_path / "b",
                   apply_to_theme(base, parse('[[rule]]\nname = "Keep"\nmatch = { text = "contents" }\naction = "keep"')))
    assert "Contents" in kept.source_text and kept.rule_counts == {"Keep": 1}


@pytest.mark.parametrize("text, problem", [
    ('[[rule]]\nname = "x"\nmatch = { page = 2 }\naction = { heading = 1 }', "can only drop"),
    ('[[rule]]\nname = "x"\nmatch = { kind = "table" }\naction = { heading = 1 }', "only paragraphs and headings"),
    ('[[rule]]\nname = "x"\nmatch = { colour = "red" }\naction = "drop"', "unknown"),
    ('[[rule]]\nname = "x"\nmatch = { kind = "paragraph" }\naction = { nested_table = "unwrap_outer" }', "kind = \"table\""),
    ('[[rule]]\nname = "x"\nmatch = { kind = "paragraph" }\naction = "drop"\nexpect = "lots"', "expect"),
    ('[settings]\nheadings = { "A" = 9 }', "heading level"),
    ('[settings]\nlua = "os.exit()"', "unknown setting"),
    ('[[rule]]\nname = "x"\nmatch = { kind = "paragraph" }\naction = "drop"\n[[rule]]\nname = "x"\nmatch = { kind = "table" }\naction = "drop"', "used twice"),
])
def test_invalid_mappings_are_refused(text, problem):
    with pytest.raises(MappingError) as e:
        parse(text)
    assert any(problem in p for p in e.value.problems), e.value.problems


def test_fingerprint_identifies_what_a_mapping_does():
    assert fingerprint(parse("version = 1")) == ""
    a = fingerprint(parse(RULES))
    assert a and a == fingerprint(parse("# comment\n" + RULES)) and a != fingerprint(parse(RULES.replace("1+", "2+")))
