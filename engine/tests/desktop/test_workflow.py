"""Proposals, the agent loop (with the deterministic fake agent in the real sandbox),
privacy and escape checks, batch treatments, export and recovery."""
from __future__ import annotations

import json
import os
import queue
import socket
import sys
import threading
import time
from pathlib import Path

import pytest

from livery.desktop import builder, conversation, events, export, jobs, patterns, proposals, review
from livery.desktop.design import Design
from livery.desktop.store import Conflict, NotFound, Project, read_json
from livery.desktop.workspace import workspace_dir

from .docs import SENTINEL, SENTINEL_NUMBER


def _wait_jobs(queued):
    for j in queued:
        try:
            jobs.wait(j, timeout=600)
        except Exception:  # noqa: BLE001 - the job records its own failure
            pass


def _turn(project, provider, text, selections, timeout=240, images=()):
    """Send one message and collect the agent events until the turn is processed."""
    q = events.subscribe(project.id)
    try:
        conv = conversation.get(project, provider)
        conv.send(text, selections, list(images))
        got = []
        deadline = time.time() + timeout
        while time.time() < deadline:
            try:
                ev = q.get(timeout=5)
            except queue.Empty:
                continue
            if ev["type"] != "agent":
                continue
            got.append(ev["event"])
            if ev["event"]["type"] == "turn" and ev["event"].get("state") == "finished":
                return got
            if ev["event"]["type"] == "error" and ev["event"].get("kind") == "invalid_design":
                return got
        raise AssertionError(f"turn did not finish: {got}")
    finally:
        events.unsubscribe(project.id, q)


def _selection(project, key, element):
    doc = project.ids[key]
    smap = read_json(builder.render_root(project, doc, 1, project.accepted()) / builder.SELECTION_FILE)
    el = next(e for e in smap["elements"] if e["id"] == element)
    return {"doc": doc, "element": element, "kind": el["kind"], "label": el["label"], "table": el.get("table")}


def _ensure_accepted_rendered(project):
    _wait_jobs([jobs.queue_render(project, d["id"], d["current"], project.accepted())
                for d in project.documents() if not builder.build_status(project, d["id"], d["current"], project.accepted())])


# ---------------------------------------------------------------- proposals
def test_proposal_lifecycle_accept_stale_discard_undo(project):
    _ensure_accepted_rendered(project)
    base_rev = project.accepted()
    base = project.design(base_rev)
    doc = project.ids["small"]

    def with_override(color):
        d = Design(base.theme_text, json.loads(json.dumps(base.design)), base.custom_text)
        d.design.setdefault("overrides", {})[f"{doc}/h0001"] = {"color": color}
        return d

    a = proposals.create(project, with_override("#8A1C1C"), base=base_rev, origin="manual", note="red")
    b = proposals.create(project, with_override("#1C8A1C"), base=base_rev, origin="manual", note="green")
    assert a["scope"]["scope"] == "elements" and a["documents"] == [doc]
    with pytest.raises(Conflict, match="render the proposal"):
        proposals.accept(project, a["id"])
    _wait_jobs(proposals.render_real(project, a["id"]))
    assert proposals.get(project, a["id"])["real"][doc]["state"] == "done"
    rev = proposals.accept(project, a["id"])["revision"]
    assert project.accepted() == rev
    # b started from the old revision: stale, refused
    _wait_jobs(proposals.render_real(project, b["id"]))
    with pytest.raises(Conflict, match="stale"):
        proposals.accept(project, b["id"])
    # untouched documents carried their build over; the changed one uses the proposal render
    for d in project.documents():
        assert builder.build_status(project, d["id"], d["current"], rev) is not None
    c = proposals.create(project, with_override("#1C1C8A"), base=rev, origin="manual", note="blue")
    proposals.discard(project, c["id"])
    assert proposals.get(project, c["id"])["status"] == "discarded"
    assert project.accepted() == rev
    src = project.source_dir(doc, 1) / "source.docx"
    digest = project.document(doc)["versions"][0]["sha256"]
    proposals.undo(project)
    assert project.accepted() == base_rev
    from livery.desktop.store import sha256_of
    assert sha256_of(src) == digest and not os.access(src, os.W_OK)


def test_remove_any_accepted_change(project):
    _ensure_accepted_rendered(project)
    small, report = project.ids["small"], project.ids["report"]

    def accept_override(ref, color, note):
        base_rev = project.accepted()
        d = project.design(base_rev)
        d = d.with_design({**d.design, "overrides": {**d.design.get("overrides", {}), ref: {"color": color}}})
        q = proposals.create(project, d, base=base_rev, origin="manual", note=note)
        _wait_jobs(proposals.render_real(project, q["id"]))
        return proposals.accept(project, q["id"])["revision"]

    a = accept_override(f"{small}/h0001", "#8A1C1C", "small red")
    b = accept_override(f"{report}/h0001", "#1C8A1C", "report green")
    assert [c["id"] for c in project.changes()][-2:] == [a, b]
    # an earlier change becomes a proposal without it; the later one stays
    q = proposals.remove_change(project, a)["proposal"]
    assert q["reverts"] == a and q["scope"]["refs"] == [f"{small}/h0001"]
    over = proposals.design_of(project, q["id"]).design["overrides"]
    assert f"{small}/h0001" not in over and over[f"{report}/h0001"] == {"color": "#1C8A1C"}
    _wait_jobs(proposals.render_real(project, q["id"]))
    r = proposals.accept(project, q["id"])["revision"]
    changes = {c["id"]: c for c in project.changes()}
    assert changes[a]["reverted_by"] == r and changes[r]["reverts"] == a
    with pytest.raises(Conflict, match="already removed"):
        proposals.remove_change(project, a)
    # a later change to the same element blocks removing the earlier one on its own
    c = accept_override(f"{report}/h0001", "#1C1C8A", "report blue")
    with pytest.raises(Conflict, match="later change also changed the override"):
        proposals.remove_change(project, b)
    # the latest change is undone at once
    assert proposals.remove_change(project, c) == {"revision": r, "proposal": None}
    assert project.accepted() == r


def test_failed_verification_blocks_accept(project):
    base_rev = project.accepted()
    base = project.design(base_rev)
    bad = Design(base.theme_text, base.design, "#let custom-rules(body) = {\n  show table: it => []\n  body\n}\n")
    meta = proposals.create(project, bad, base=base_rev, origin="manual", note="drops tables")
    assert meta["scope"]["scope"] == "global"
    _wait_jobs(proposals.render_real(project, meta["id"], [project.ids["small"]]))
    with pytest.raises(Conflict):
        proposals.accept(project, meta["id"])
    proposals.discard(project, meta["id"])


def test_done_documents_are_flagged_only_when_their_pages_change(project):
    _ensure_accepted_rendered(project)
    small, report = project.ids["small"], project.ids["report"]

    def proposal(ref, color):
        base_rev = project.accepted()
        d = project.design(base_rev)
        d = d.with_design({**d.design, "overrides": {**d.design.get("overrides", {}), ref: {"color": color}}})
        q = proposals.create(project, d, base=base_rev, origin="manual", note=f"{ref} {color}")
        _wait_jobs(proposals.render_real(project, q["id"]))
        return q["id"]

    def state(doc):
        st = project.state()
        return review.state_of(project, st["documents"][doc], st["accepted"], st["history"])

    try:
        review.mark(project, small, True)
        assert state(small)["state"] == "done" and state(report)["state"] == "open"
        # a change to another document leaves it done
        q = proposal(f"{report}/h0001", "#2B5F8A")
        assert small not in proposals.get(project, q)["done_impact"]
        proposals.accept(project, q)
        assert state(small)["state"] == "done"
        # a change to it is announced on the proposal, and flags it once accepted
        q = proposal(f"{small}/h0001", "#8A1C1C")
        impact = proposals.get(project, q)["done_impact"][small]
        assert impact["state"] == "changed" and impact["pages"]
        rev = proposals.accept(project, q)["revision"]
        st = state(small)
        assert st["state"] == "recheck" and st["reason"] == "design" and st["pages"] == impact["pages"]
        assert [c["id"] for c in st["changes"]][-1] == rev
        stats = review.stats(project)
        assert stats["review"]["recheck"] >= 1 and any(r["doc"] == small for r in stats["recheck"])
        # approving it again with the new design clears the flag
        review.mark(project, small, True)
        assert state(small)["state"] == "done"
        # undoing the change it was approved with flags it again
        proposals.undo(project)
        assert state(small)["state"] == "recheck"
    finally:
        review.mark(project, small, False)


def test_marking_done_needs_a_passing_render(project):
    _ensure_accepted_rendered(project)
    with pytest.raises(NotFound):
        review.mark(project, "dnothere", True)
    stats = review.stats(project)
    assert stats["documents"] == len(project.documents())
    assert sum(stats["verification"].values()) == stats["documents"]


# ---------------------------------------------------------------- agent loop
def test_fake_agent_edits_a_table_through_a_proposal(project):
    _ensure_accepted_rendered(project)
    sel = [_selection(project, "report", "t0006")]
    got = _turn(project, "fake", "Make this header #AA3311 please", sel)
    kinds = [e["type"] for e in got]
    assert "text_delta" in kinds and "tool" in kinds and "done" in kinds
    pid = got[-1]["proposal"]
    assert pid, got
    meta = proposals.get(project, pid)
    design = proposals.design_of(project, pid).design
    ref = f"{project.ids['report']}/t0006"
    assert design["assignments"][ref] == "Fake table"
    assert design["treatments"]["Fake table"]["style"]["header_fill"] == "#AA3311"
    synthetic = read_json(project.proposal_dir(pid) / "synthetic.json")
    assert synthetic["ok"], synthetic
    assert (project.proposal_dir(pid) / "synthetic").glob("doc-*.pdf")
    # the agent asked for a preview through render-preview and got an answer
    conv = conversation.active(project)
    assert list((conv.ws / ".responses").glob("*.json"))
    assert list((conv.ws / "previews").glob("doc-01/page-*.png"))
    _wait_jobs(proposals.render_real(project, pid))
    rev = proposals.accept(project, pid)["revision"]
    assert project.design(rev).design["assignments"][ref] == "Fake table"
    assert meta["focus"] == [project.ids["report"]]


def _photo(width=3000, height=1500):
    """A JPEG with EXIF (camera and author) as a phone or a screenshot tool would save it."""
    import io
    from PIL import Image
    im = Image.new("RGB", (width, height), "#1F3A5F")
    exif = Image.Exif()
    exif[0x010F] = "PhoneMaker"  # Make
    exif[0x013B] = SENTINEL      # Artist
    out = io.BytesIO()
    im.save(out, format="JPEG", exif=exif)
    return out.getvalue()


def test_reference_images_are_reencoded_without_metadata():
    import io
    from PIL import Image
    from livery.desktop.workspace import normalize_reference
    png = normalize_reference(_photo())
    assert png.startswith(b"\x89PNG") and SENTINEL.encode() not in png and b"PhoneMaker" not in png
    with Image.open(io.BytesIO(png)) as im:
        assert im.size == (2000, 1000) and not im.getexif()
    for bad in (b"not an image", b"%PDF-1.7 ...", b"GIF89a"):
        with pytest.raises(ValueError):
            normalize_reference(bad)


def test_reference_images_reach_the_agent(project):
    """An attached image lands in the workspace as a clean PNG, is named in the prompt and
    can be read by the agent in its sandbox; later images are numbered on."""
    from livery.desktop.workspace import normalize_reference
    _ensure_accepted_rendered(project)
    sel = [_selection(project, "report", "t0006")]
    got = _turn(project, "fake", "Make all tables look like this #224466", sel, images=[normalize_reference(_photo())])
    said = " ".join(e["text"] for e in got if e["type"] == "text")
    assert "Looked at references/ref-01.png:png" in said, said
    got = _turn(project, "fake", "And this one", sel, images=[normalize_reference(_photo(400, 300))])
    said = " ".join(e["text"] for e in got if e["type"] == "text")
    assert "references/ref-02.png:png" in said, said
    ws = conversation.active(project).ws
    assert sorted(f.name for f in (ws / "references").iterdir()) == ["ref-01.png", "ref-02.png"]
    assert SENTINEL.encode() not in (ws / "references" / "ref-01.png").read_bytes()
    hist = conversation.history(project, "fake")
    assert [e.get("attachments") for e in hist if e["role"] == "user"][-1] == [sel[0]["label"], "Image 2"]


def test_private_content_never_reaches_the_agent(project):
    """The sentinel words of the report never appear in the workspace, the prompts the
    agent received, its tool responses or the agent-side records."""
    conv = conversation.get(project, "fake")
    _turn(project, "fake", "Make this header #123456", [_selection(project, "report", "t0006")])
    ws = conv.ws
    hits = []
    for path in [ws, *Path(ws).parent.glob(".render-*")]:
        for f in Path(path).rglob("*"):
            if f.is_file() and not f.is_symlink():
                data = f.read_bytes()
                if SENTINEL.encode() in data or SENTINEL_NUMBER.encode() in data:
                    hits.append(str(f))
    assert not hits, hits
    assert (ws / ".fake-log").exists()
    # sanity: the sentinel really is in the private side
    src = builder.inventory(project, project.ids["report"], 1)
    assert any(SENTINEL in e["text"] for e in src)


def test_sandbox_escapes_are_blocked(project):
    srv = socket.socket()
    srv.bind(("127.0.0.1", 0))
    srv.listen(1)
    port = srv.getsockname()[1]
    threading.Thread(target=lambda: srv.accept(), daemon=True).start()
    home = Path.home()
    private = [str(project.root), str(project.source_dir(project.ids["report"], 1) / "source.docx"),
               str(Path(os.environ["LIVERY_HOME"]) / "projects"), str(home / "Documents"),
               str(home / ".claude"), str(home / ".codex"), str(home / ".ssh"), str(home / ".config")]
    private += [str(home)]
    prompt = "PROBE " + " ".join(private) + f" PORT={port} LINK={private[1]}"
    got = _turn(project, "fake", prompt, [])
    text = next(e["text"] for e in got if e["type"] == "text" and e["text"].startswith("PROBE-RESULT"))
    report = json.loads(text.split(" ", 1)[1])
    for p in private[:-1]:
        assert report[p].startswith("blocked"), (p, report[p])
    # The home folder shows only the parents of the read-only runtime mounts (the
    # interpreter and the agent binary), never the user's files.
    listed = report[str(home)].removeprefix("listed:").split(",") if report[str(home)].startswith("listed:") else []
    runtime = {".local", ".superset", ".cache", ".npm-global"}
    for mount in (Path(sys.prefix).resolve(), Path(sys.executable).resolve()):  # e.g. ~/work/<repo>/.venv on CI
        if mount.is_relative_to(home):
            runtime.add(mount.relative_to(home).parts[0])
    assert set(listed) <= runtime, listed
    assert report[f"tcp:{port}"].startswith("blocked")
    assert report["internet"].startswith("blocked")
    assert report["/proc/1/environ"].startswith("blocked") or "LIVERY" not in report["/proc/1/environ"]
    assert report["/run"].startswith("blocked")
    assert not any(k.startswith(("LIVERY", "SSH_", "DBUS", "XDG", "DISPLAY", "WAYLAND")) for k in report["env"])
    assert set(report["env"]) <= {"PATH", "HOME", "LANG", "TERM", "USER", "FAKE_PROFILE", "PWD", "SHLVL", "_", "OLDPWD"}
    # the planted symlink is refused instead of followed
    assert report["symlink"] == "planted"
    assert any(e["type"] == "error" and e.get("kind") == "invalid_design" and "regular file" in e["message"]
               for e in got), got
    srv.close()


def test_cancel_and_recoverable_errors(project):
    conv = conversation.get(project, "fake")
    q = events.subscribe(project.id)
    try:
        conv.send("SLEEP for a while", [])
        time.sleep(1.5)
        assert conv.session.busy
        conv.cancel()
        deadline = time.time() + 30
        err = None
        while time.time() < deadline and err is None:
            try:
                ev = q.get(timeout=2)
            except queue.Empty:
                continue
            if ev["type"] == "agent" and ev["event"]["type"] == "error":
                err = ev["event"]
        assert err and err["kind"] == "cancelled" and err["recoverable"]
    finally:
        events.unsubscribe(project.id, q)
    time.sleep(1)
    got = _turn(project, "fake", "FAIL_AUTH please", [])
    errs = [e for e in got if e["type"] == "error"]
    assert errs and errs[0]["kind"] == "auth"


def test_switching_providers_keeps_conversations_separate(project):
    fake = conversation.get(project, "fake")
    key, agent_session = fake.rec["key"], fake.rec["agent_session"]
    other = conversation.get(project, "claude")  # switching away closes the fake session
    assert fake.session is None and other.provider == "claude"
    back = conversation.get(project, "fake")
    assert back is not fake
    if back.rec is not None:  # restored: same workspace and agent session
        assert back.rec["key"] == key and back.rec["agent_session"] == agent_session
    hist = conversation.history(project, "fake")
    assert any(h["role"] == "user" for h in hist)
    assert conversation.history(project, "claude") == []


def test_clearing_starts_a_fresh_conversation(project):
    conv = conversation.get(project, "fake")
    conv.ensure([])
    key = conv.rec["key"]
    assert conversation.history(project, "fake")
    conversation.clear(project, "fake")
    assert conv.session is None and conversation.active(project) is None
    assert conversation.history(project, "fake") == []
    assert list((project.root / "chats" / "archive").glob("fake-*.jsonl"))
    fresh = conversation.get(project, "fake")
    assert fresh.rec["key"] == key  # same workspace
    assert fresh.rec["agent_session"] is None and fresh.rec["needs_summary"]


# ---------------------------------------------------------------- batch & export
def test_table_signatures_find_reviewed_matches(project):
    _ensure_accepted_rendered(project)
    found = patterns.candidates(project, project.ids["report"], "t0006")
    by_doc = {(c["doc"], c["element"]): c for c in found["candidates"]}
    assert (project.ids["report"], "t0006") in by_doc
    small = [c for c in found["candidates"] if c["doc"] == project.ids["small"]]
    assert any(c["match"] == "match" for c in small), found
    assert not any(c["doc"] == project.ids["plain"] and c["match"] == "match" for c in found["candidates"])
    base = project.design(project.accepted())
    if "Data table" not in base.design.get("treatments", {}):
        d = Design(base.theme_text, json.loads(json.dumps(base.design)), base.custom_text)
        d.design.setdefault("treatments", {})["Data table"] = {"kind": "table", "style": {"header_fill": "#22446B",
                                                                                          "header_color": "#FFFFFF"}}
        meta = proposals.create(project, d, base=project.accepted(), origin="manual", note="treatment")
        _wait_jobs(proposals.render_real(project, meta["id"]))
        proposals.accept(project, meta["id"])
    refs = [{"doc": c["doc"], "element": c["element"]} for c in found["candidates"] if c["match"] == "match"]
    meta = patterns.apply(project, "Data table", refs)
    assert set(meta["documents"]) >= {project.ids["report"], project.ids["small"]}
    _wait_jobs(proposals.render_real(project, meta["id"]))
    rev = proposals.accept(project, meta["id"])["revision"]
    assignments = project.design(rev).design["assignments"]
    assert all(assignments[f"{r['doc']}/{r['element']}"] == "Data table" for r in refs)
    with pytest.raises(Conflict):
        patterns.apply(project, "Data table", [{"doc": project.ids["report"], "element": "h0001"}])


def test_export_keeps_going_past_a_failed_document(desktop_home, doc_files, tmp_path):
    p = Project.create("Export")
    good, gv = p.add_source(doc_files["small"])
    bad, bv = p.add_source(doc_files["broken"])
    _wait_jobs([jobs.queue_import(p, good, gv), jobs.queue_import(p, bad, bv)])
    assert p.document(bad)["versions"][0]["status"] == "failed"
    result = export.export(p, tmp_path / "out")
    rows = {r["doc"]: r for r in result["documents"]}
    assert rows[good]["exported"], (rows[good]["failures"], rows[good]["warnings"])
    assert rows[good]["exported"] and (tmp_path / "out" / rows[good]["pdf"]).exists()
    assert not rows[bad]["exported"] and rows[bad]["failures"]
    assert (tmp_path / "out" / "livery-report.md").exists()
    assert result["summary"] == {"exported": 1, "failed": 1, "warnings": result["summary"]["warnings"], "total": 2}
    one = export.export(p, tmp_path / "one", only=[good])
    assert [r["doc"] for r in one["documents"]] == [good] and one["summary"]["exported"] == 1


def test_source_problems_block_export_in_every_revision(desktop_home, tmp_path):
    """Unresolved tracked changes FAIL at import; that holds for every design revision."""
    from tests.conftest import FIXTURES
    p = Project.create("Tracked")
    doc, v = p.add_source(FIXTURES / "tracked.docx")
    builder.import_document(p, doc, v)
    st = builder.render_document(p, doc, v, "r0001")
    assert st["report_status"] == "FAIL"
    report = builder.report_of(builder.render_root(p, doc, v, "r0001"))
    assert any(c["name"] == "tracked_changes" and c["status"] == "FAIL" for c in report["checks"])
    result = export.export(p, tmp_path / "out")
    assert result["summary"]["exported"] == 0 and not list((tmp_path / "out").glob("*.pdf"))


def test_restart_recovers_interrupted_work(desktop_home, doc_files):
    p = Project.create("Recover")
    doc, version = p.add_source(doc_files["plain"])
    p.set_version_status(doc, version, status="importing")  # the app died mid-import
    stale = builder.render_root(p, doc, version, "r0001").with_name(".r0001.tmp")
    stale.mkdir(parents=True)
    notes = jobs.recover(Project.open(p.id))
    assert any("re-importing" in n for n in notes) and not stale.exists()
    deadline = time.time() + 300
    while time.time() < deadline and p.document(doc)["versions"][0]["status"] != "imported":
        time.sleep(1)
    assert p.document(doc)["versions"][0]["status"] == "imported"
    for _ in range(300):
        st = builder.build_status(p, doc, version, "r0001")
        if st and st["state"] == "done":
            break
        time.sleep(1)
    assert builder.build_status(p, doc, version, "r0001")["state"] == "done"


def test_twin_ids_are_aligned_to_the_real_document(desktop_home, tmp_path):
    """A heading dropped by its text exists only in the lorem twin; the workspace must still
    address every real element by its real ID."""
    from livery.desktop import workspace

    from .docs import contents_doc
    p = Project.create("Contents")
    doc, v = p.add_source(contents_doc(tmp_path / "review.docx"))
    builder.import_document(p, doc, v)
    real = builder.inventory(p, doc, v)
    twin = builder.inventory(p, doc, v, twin=True)
    assert len(twin) == len(real) + 1  # the premise: the twin kept the "Contents" heading
    info = workspace.create(p, "s-align", [doc])
    sample = Path(info["dir"]) / "samples" / info["docs"][doc]
    ws_inv = read_json(sample / "elements.json")
    real_ids = {e["id"]: e["kind"] for e in real}
    mapped = [e for e in ws_inv if not e["id"].startswith("x")]
    assert {e["id"]: e["kind"] for e in mapped} == real_ids
    table = next(e["id"] for e in real if e["kind"] == "table")
    assert f'#metadata("{table}") <we-s>' in (sample / "content.typ").read_text()


def test_workspace_lives_outside_projects(project):
    conv = conversation.get(project, "fake")
    if conv.rec is None:
        conv.ensure([])
    ws = Path(conv.rec["dir"]).resolve()
    assert ws == workspace_dir(project, conv.rec["key"]).resolve()
    assert not str(ws).startswith(str(project.root.resolve()))


# ---------------------------------------------------------------- import mapping
MAPPING = """version = 1

[[rule]]
name = "Leftover page"
match = { page = 2 }
action = "drop"
expect = "1+"

[[rule]]
name = "Bold numbered titles"
match = { kind = "paragraph", bold = true, numbered = true }
action = { heading = "from_number" }

[[rule]]
name = "Container tables"
match = { kind = "table", contains_table = true }
action = { nested_table = "unwrap_outer" }
"""


def _mapping_project(tmp_path, name):
    from . import docs
    p = Project.create(name)
    doc, version = p.add_source(docs.mapping_doc(tmp_path / f"{name}.docx"))
    assert builder.import_document(p, doc, version)["status"] in ("PASS", "WARN")
    inv = builder.inventory(p, doc, version)
    overview = next(e["id"] for e in inv if e["kind"] == "heading" and e["text"] == "Overview")
    scope = next(e["id"] for e in inv if e["kind"] == "paragraph" and e["text"].startswith("1.1"))
    base = p.design("r0001")
    colours = base.with_design({**base.design, "overrides": {f"{doc}/{overview}": {"color": "#8A1C1C"},
                                                           f"{doc}/{scope}": {"color": "#1C1C8A"}}})
    q = proposals.create(p, colours, base="r0001", origin="manual", note="colours")
    _wait_jobs(proposals.render_real(p, q["id"]))
    rev = proposals.accept(p, q["id"])["revision"]
    _wait_jobs([jobs.queue_render(p, doc, version, rev)])
    return p, doc, scope


def _check_mapped(p, doc, rev, scope):
    """The accepted revision rev converts with the mapping: page 2 gone, the bold title a
    heading, the container table unwrapped; the heading kept its colour under its new ID and
    the paragraph that became a heading lost its paragraph colour."""
    assert p.accepted() == rev and "Leftover page" in p.design(rev).mapping_text
    inv = builder.inventory(p, doc, 1)
    headings = [e["text"] for e in inv if e["kind"] == "heading"]
    assert headings == ["Overview", "1.1 Scope"], headings
    assert not any("Distribution" in e["text"] for e in inv)
    assert sum(e["kind"] == "table" for e in inv) == 1
    overview = next(e["id"] for e in inv if e["text"] == "Overview")
    assert p.design(rev).design["overrides"] == {f"{doc}/{overview}": {"color": "#8A1C1C"}}
    report = builder.report_of(builder.render_root(p, doc, 1, rev))
    checks = {c["name"]: c for c in report["checks"]}
    assert checks["mapping_text"]["status"] == "PASS", checks["mapping_text"]
    assert checks["mapping_rules"]["status"] == "PASS", checks["mapping_rules"]
    assert any(SENTINEL in d for d in checks["mapping_rules"]["details"])  # the removal is on the record, locally
    assert report["status"] != "FAIL", [c for c in report["checks"] if c["status"] == "FAIL"]


def test_mapping_proposal_carries_references_to_the_new_conversion(desktop_home, tmp_path):
    p, doc, scope = _mapping_project(tmp_path, "Direct")
    base_rev = p.accepted()
    base = p.design(base_rev)
    import dataclasses
    meta = proposals.create(p, dataclasses.replace(base, mapping_text=MAPPING), base=base_rev, origin="manual",
                            note="mapping")
    assert meta["scope"]["mapping"] and meta["scope"]["scope"] == "global" and meta["remap"]["docs"] == [doc]
    with pytest.raises(Conflict, match="render"):
        proposals.accept(p, meta["id"])
    _wait_jobs(proposals.render_real(p, meta["id"]))
    assert proposals.get(p, meta["id"])["remap"]["dropped"] == [f"{doc}/{scope}"]
    rev = proposals.accept(p, meta["id"])["revision"]
    _check_mapped(p, doc, rev, scope)
    # undo goes back to the default conversion and its IDs
    proposals.undo(p)
    assert any(e["text"].startswith("1.1") and e["kind"] == "paragraph" for e in builder.inventory(p, doc, 1))


def test_agent_changes_the_mapping_through_a_proposal(desktop_home, tmp_path):
    p, doc, scope = _mapping_project(tmp_path, "Agent")
    got = _turn(p, "fake", "MAPPING\n" + MAPPING, [])
    text = "".join(e.get("text", "") for e in got if e["type"] == "text_delta")
    preview = json.loads(text.split("PREVIEW-RESULT ", 1)[1].split("\n")[0])
    assert preview["ok"], preview
    assert preview["docs"]["doc-01"]["mapping_rules"] == {"Leftover page": 2, "Bold numbered titles": 1,
                                                          "Container tables": 1}
    assert any("could not be carried" in n and f"doc-01/{scope}" in n for n in preview["notes"])
    ws = conversation.active(p).ws
    sample = read_json(ws / "samples" / "doc-01" / "elements.json")
    assert [e["kind"] for e in sample].count("heading") == 2
    styles = read_json(ws / "samples" / "doc-01" / "styles.json")
    assert styles["pages"] == 3 and any(s["style"] == "Normal" and 2 in s["pages"] for s in styles["styles"])
    hits = [str(f) for f in Path(ws).rglob("*") if f.is_file() and not f.is_symlink()
            and (SENTINEL.encode() in f.read_bytes() or SENTINEL_NUMBER.encode() in f.read_bytes())]
    assert not hits, hits
    pid = got[-1]["proposal"]
    meta = proposals.get(p, pid)
    assert meta["scope"]["mapping"] and meta["remap"]["docs"] == []  # the only document was in the workspace
    _wait_jobs(proposals.render_real(p, pid))
    rev = proposals.accept(p, pid)["revision"]
    _check_mapped(p, doc, rev, scope)


def test_private_custom_style_names_are_hidden(tmp_path):
    from docx import Document
    from livery.desktop.workspace import style_usage
    doc = Document()
    doc.styles.add_style(f"{SENTINEL} Title", 1)
    doc.add_paragraph("Lorem ipsum", style=f"{SENTINEL} Title")
    doc.add_paragraph("Dolor sit")
    doc.save(tmp_path / "t.docx")
    names = [s["style"] for s in style_usage(tmp_path / "t.docx", {SENTINEL.lower(), "normal"})["styles"]]
    assert names == ["(hidden style 1)", "Normal"]
