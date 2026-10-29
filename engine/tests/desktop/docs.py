"""Word documents for the desktop tests, generated with python-docx.

Each covers a structure the plan names: repeated text, merged (grouped) table headers, a
long table that crosses pages, a paragraph that spans pages, an image, structures the
selection does not support (lists, a nested table), and a sentinel phrase that must never
reach an agent.
"""
from __future__ import annotations

import io
from pathlib import Path

from docx import Document
from docx.shared import Inches
from PIL import Image

SENTINEL = "Zorblaxquint"  # appears in private documents only; tests scan for it
SENTINEL_NUMBER = "8675309421"


def _png() -> io.BytesIO:
    buf = io.BytesIO()
    Image.new("RGB", (320, 160), (40, 90, 160)).save(buf, "PNG")
    buf.seek(0)
    return buf


def _grouped_table(doc, rows: int, label: str):
    """Two header rows: a merged group cell over columns 2-3."""
    t = doc.add_table(rows=2 + rows, cols=4)
    t.style = "Table Grid"
    h0, h1 = t.rows[0].cells, t.rows[1].cells
    h0[0].text = "Item"
    h0[1].merge(h0[2]).text = "Quarterly figures"
    h0[3].text = "Notes"
    h1[0].text = ""
    h1[1].text = "Q1"
    h1[2].text = "Q2"
    h1[3].text = ""
    h0[0].merge(h1[0])
    h0[3].merge(h1[3])
    for i in range(rows):
        cells = t.rows[2 + i].cells
        cells[0].text = f"{label} row {i + 1}"
        cells[1].text = str(100 + i)
        cells[2].text = str(200 + i)
        cells[3].text = "Repeated note"
    return t


def _simple_table(doc, rows: int, cols: int = 3, label: str = "Line"):
    t = doc.add_table(rows=rows + 1, cols=cols)
    t.style = "Table Grid"
    for c in range(cols):
        t.rows[0].cells[c].text = f"Column {c + 1}"
    for r in range(rows):
        for c in range(cols):
            t.rows[r + 1].cells[c].text = f"{label} {r + 1}.{c + 1}"
    return t


def report_doc(path: Path, *, sentinel: bool = True, title: str = "Quarterly Operations Report") -> Path:
    doc = Document()
    doc.core_properties.title = title
    doc.add_heading("Overview", level=1)
    p = "This paragraph repeats. " * 3
    doc.add_paragraph(p)
    doc.add_paragraph(p)  # repeated text: identical paragraphs must keep distinct IDs
    if sentinel:
        doc.add_paragraph(f"Confidential reference {SENTINEL} with account {SENTINEL_NUMBER}.")
    doc.add_heading("Figures by quarter", level=2)
    _grouped_table(doc, 4, "Region")
    doc.add_paragraph("A short paragraph after the grouped table.")
    doc.add_heading("Detailed ledger", level=2)
    _simple_table(doc, 70, label="Entry")  # long table: crosses pages
    doc.add_heading("Narrative", level=2)
    doc.add_paragraph(" ".join(f"Sentence number {i} explains a detail of the operation." for i in range(260)))
    doc.add_paragraph("First list item", style="List Bullet")
    doc.add_paragraph("Second list item", style="List Bullet")
    doc.add_picture(_png(), width=Inches(3))
    outer = doc.add_table(rows=1, cols=2)
    outer.style = "Table Grid"
    outer.rows[0].cells[0].text = "Outer cell"
    inner = outer.rows[0].cells[1].add_table(rows=1, cols=1)
    inner.rows[0].cells[0].text = "Nested cell"
    doc.add_paragraph("Closing paragraph.")
    doc.save(path)
    return path


def contents_doc(path: Path) -> Path:
    """A "Contents" heading the converter drops by its text: the lorem twin keeps it, so the
    twin has one element more than the real document."""
    doc = Document()
    doc.core_properties.title = "Annual review"
    doc.add_heading("Contents", level=1)
    doc.add_paragraph("Introduction comes first.")
    doc.add_heading("Introduction", level=1)
    doc.add_paragraph("The year in short.")
    _simple_table(doc, 3, label="Item")
    doc.save(path)
    return path


def toc_table_doc(path: Path) -> Path:
    """A hand-made contents table: cells that start with a typed number ("1. Overview")."""
    doc = Document()
    doc.core_properties.title = "Plan"
    t = doc.add_table(rows=3, cols=2)
    t.style = "Table Grid"
    for i, (name, page) in enumerate([("1. Overview", "3"), ("2. Market and Audience", "4"), ("3. Risks", "9")]):
        t.rows[i].cells[0].text = name
        t.rows[i].cells[1].text = page
    doc.add_paragraph("Body text.")
    doc.save(path)
    return path


def small_doc(path: Path, *, title: str = "Supplier Summary", grouped: bool = True) -> Path:
    doc = Document()
    doc.core_properties.title = title
    doc.add_heading("Suppliers", level=1)
    doc.add_paragraph("Suppliers are listed by quarter below.")
    if grouped:
        _grouped_table(doc, 3, "Supplier")
    _simple_table(doc, 2, cols=2, label="Contact")
    doc.save(path)
    return path


def broken_doc(path: Path) -> Path:
    path.write_bytes(b"PK\x03\x04 this is not really a docx")
    return path


def mapping_doc(path: Path) -> Path:
    """What a project's mapping fixes: a leftover page after the cover, titles typed as bold
    numbered Normal paragraphs, and a table used as a container for another table. The
    sentinel is on the leftover page and in the body."""
    from docx.enum.text import WD_BREAK

    doc = Document()
    doc.core_properties.title = "Operating Plan"
    doc.add_paragraph("Operating Plan")
    doc.add_paragraph("Prepared for the board").runs[0].add_break(WD_BREAK.PAGE)
    doc.add_paragraph(f"Distribution list {SENTINEL}")
    doc.add_paragraph("Internal use only").runs[0].add_break(WD_BREAK.PAGE)
    doc.add_heading("Overview", level=1)
    doc.add_paragraph("The plan in short.")
    doc.add_paragraph().add_run("1.1 Scope").bold = True
    doc.add_paragraph(f"What the plan covers, account {SENTINEL_NUMBER}.")
    outer = doc.add_table(rows=1, cols=1)
    outer.style = "Table Grid"
    cell = outer.rows[0].cells[0]
    cell.paragraphs[0].text = "Figures"
    inner = cell.add_table(rows=2, cols=2)
    for r in range(2):
        for c in range(2):
            inner.rows[r].cells[c].text = f"Value {r}{c}"
    doc.add_paragraph("Closing words.")
    doc.save(path)
    return path
