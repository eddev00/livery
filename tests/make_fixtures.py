"""Generate the .docx fixtures used by the test-suite.

Run:  uv run python tests/make_fixtures.py
Requires pandoc on PATH (or in .tools/bin) and the dev dependencies
(python-docx, pillow).
"""
from __future__ import annotations

import os
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
FIX = ROOT / "tests" / "fixtures"
TOOLS = ROOT / ".tools" / "bin"
os.environ["PATH"] = f"{TOOLS}{os.pathsep}{os.environ['PATH']}"

FEATURES_MD = r"""---
title: Feature Coverage Document
author: Fixture Bot
date: 2026-09-14
---

::: {custom-style="Chapter Title"}
Introduction to the System
:::

This paragraph contains characters that are special in Typst: # hash, *stars*,
_underscores_, @at, \<angle\> brackets, $dollar, a back\\slash, a `tilde~` and
[square] brackets. It also has "smart quotes", 'single ones', an en–dash, an
em—dash, an ellipsis… and a non-breaking space before this: 42 kg.

Some **bold text**, some *italic text*, some ***both***, some ~~struck~~ text,
super^script^ and sub~script~, and `inline code()`.

## Lists and Nesting

- First bullet
- Second bullet with **bold**
  - Nested bullet A
  - Nested bullet B
- Third bullet

1. Step one
2. Step two
   1. Sub-step two point one
   2. Sub-step two point two
3. Step three

### A Third-Level Heading

Here is a link to [the Typst website](https://typst.app/docs) and a footnote.[^1]
Another sentence follows the footnote reference.

[^1]: This is the footnote text, which must survive verbatim.

> This is a block quote. It should be restyled but its words must not change.

```
def code_block():
    return "verbatim"
```

# Data Tables

| Name | Quantity | Unit price | Notes |
|:-----|---------:|-----------:|:------|
| Widget | 12 | 3.50 | Standard |
| Gadget | 7 | 12.00 | *Fragile* |
| Gizmo | 1 | 99.99 | Contains a # and a $ |

Table: Inventory snapshot

## Images

![A generated test image](tests/fixtures/_build/test-image.png)

::: {custom-style="Chapter Title"}
Second Chapter
:::

Final paragraph of the document, the end.
"""

TRACKED_MD = r"""---
title: Tracked Changes Document
---

# Revisions

This sentence is unchanged. [This text was inserted]{.insertion author="Reviewer" date="2026-09-01T10:00:00Z"} and
[this text was deleted]{.deletion author="Reviewer" date="2026-09-01T10:00:00Z"} from the paragraph.

[A commented sentence.]{.comment-start id="0" author="Reviewer" date="2026-09-01T10:00:00Z"}[]{.comment-end id="0"}
"""


NUMBERED_MD = r"""---
title: Numbered Headings
---

# Preface

Front matter before any numbered heading.

# 1. Alpha chapter

Text of alpha.

# 1.1 Alpha first section

All headings are Word "Heading 1" here; the typed numbers define the structure.

# 1.2. Alpha second section

More text.

# 1.2.1 Deep one

Deep text.

# 3. Gamma chapter (typed 3, should be 2)

The typed number skips 2 on purpose.
"""


def make_image(path: Path) -> None:
    from PIL import Image, ImageDraw

    img = Image.new("RGB", (400, 200), (30, 90, 160))
    d = ImageDraw.Draw(img)
    d.rectangle([20, 20, 380, 180], outline=(255, 255, 255), width=4)
    d.text((40, 90), "wordeasy fixture image", fill=(255, 255, 255))
    img.save(path)


def pandoc_md_to_docx(md: str, out: Path) -> None:
    subprocess.run(
        ["pandoc", "-f", "markdown", "-t", "docx", "--resource-path", str(ROOT), "-o", str(out)],
        input=md.encode(),
        check=True,
        cwd=ROOT,
    )


def styles_add(doc, name: str):
    from docx.enum.style import WD_STYLE_TYPE

    try:
        return doc.styles[name]
    except KeyError:
        return doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)


def make_pydocx(out: Path) -> None:
    from docx import Document
    from docx.enum.text import WD_BREAK
    from docx.shared import Pt

    doc = Document()
    sec = doc.sections[0]
    sec.header.paragraphs[0].text = "Confidential running header"
    sec.footer.paragraphs[0].text = "Footer text from Word"

    from docx.enum.style import WD_STYLE_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    toc_h = styles_add(doc, "TOC Heading")
    doc.add_paragraph("Table of Contents", style=toc_h)
    toc1 = styles_add(doc, "toc 1")
    doc.add_paragraph("1. Merged Cells and Layout\t3", style=toc1)
    doc.add_paragraph("2. Long table\t5", style=toc1)

    doc.add_heading("Merged Cells and Layout", level=1)
    # internal hyperlink (bookmark-style anchor) whose text must survive as plain text
    para = doc.add_paragraph("See ")
    h = OxmlElement("w:hyperlink")
    h.set(qn("w:anchor"), "_Toc000001")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "the long table section"
    r.append(t)
    h.append(r)
    para._p.append(h)
    para.add_run(" for details.")
    doc.add_paragraph("A paragraph followed by three empty paragraphs.")
    doc.add_paragraph("")
    doc.add_paragraph("")
    doc.add_paragraph("")
    p = doc.add_paragraph("Tab\tseparated\twords and a hard line")
    p.add_run().add_break(WD_BREAK.LINE)
    p.add_run("break inside the paragraph.")

    doc.add_heading("Merged table", level=2)
    t = doc.add_table(rows=4, cols=3)
    t.style = "Table Grid"
    hdr = t.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text = "Region", "Q1", "Q2"
    t.cell(1, 0).merge(t.cell(2, 0))
    t.cell(1, 0).text = "North (merged)"
    t.cell(1, 1).text, t.cell(1, 2).text = "100", "110"
    t.cell(2, 1).text, t.cell(2, 2).text = "120", "130"
    t.cell(3, 0).merge(t.cell(3, 2))
    t.cell(3, 0).text = "Total row spanning all columns"
    for cell, color in ((hdr[0], "1F4E79"), (hdr[1], "1F4E79"), (hdr[2], "1F4E79"), (t.cell(2, 2), "FFF2CC")):
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), color)
        cell._tc.get_or_add_tcPr().append(shd)

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_heading("Long table", level=1)
    doc.add_paragraph("The table below spans several pages.")
    t2 = doc.add_table(rows=1, cols=4)
    t2.style = "Table Grid"
    for i, h in enumerate(["ID", "Item", "Amount", "Comment"]):
        t2.rows[0].cells[i].text = h
    for i in range(1, 81):
        r = t2.add_row().cells
        r[0].text = str(i)
        r[1].text = f"Item number {i}"
        r[2].text = f"{i * 1.25:.2f}"
        r[3].text = "even row" if i % 2 == 0 else "odd row"

    doc.add_heading("Custom Style Paragraphs", level=1)
    styles = doc.styles
    from docx.enum.style import WD_STYLE_TYPE

    st = styles.add_style("Chapter Title", WD_STYLE_TYPE.PARAGRAPH)
    st.font.size = Pt(24)
    doc.add_paragraph("Styled Chapter Heading", style="Chapter Title")
    doc.add_paragraph("Body under the custom-styled chapter heading.")
    doc.save(out)


def main() -> None:
    build = FIX / "_build"
    build.mkdir(parents=True, exist_ok=True)
    make_image(build / "test-image.png")
    pandoc_md_to_docx(FEATURES_MD, FIX / "features.docx")
    pandoc_md_to_docx(TRACKED_MD, FIX / "tracked.docx")
    pandoc_md_to_docx(NUMBERED_MD, FIX / "numbered.docx")
    make_pydocx(FIX / "layout.docx")
    for f in sorted(FIX.glob("*.docx")):
        print(f"wrote {f.relative_to(ROOT)} ({f.stat().st_size} bytes)")


if __name__ == "__main__":
    sys.exit(main())
